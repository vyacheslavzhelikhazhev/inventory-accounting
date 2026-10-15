import unittest,tempfile,json,shutil,zipfile,sqlite3,contextlib,hashlib,sys,io
from pathlib import Path
from docx import Document
from store import Store,FIELDS
from enhancements import without_organization,paragraphs,replace_text
from bitrix_disk import BitrixDisk,protect

RESOURCES=Path(__file__).parent
TEST_ROOT=RESOURCES/'test'

class ProjectTests(unittest.TestCase):
    def setUp(self):
        TEST_ROOT.mkdir(parents=True,exist_ok=True)
        self.temp=tempfile.TemporaryDirectory(dir=TEST_ROOT);self.addCleanup(self.temp.cleanup)
        root=Path(self.temp.name);res=root/'res';res.mkdir()
        shutil.copyfile(RESOURCES/'Шаблон_акта.docx',res/'Шаблон_акта.docx')
        assets=[]
        for id in range(1,5):
            a={k:'' for k in FIELDS};a.update(id=id,name='Тест '+str(id),type='Ноутбук',serial='SN-'+str(id),status='На складе',location='Склад',kit='Кабель',condition='Б/у')
            if id==3:a.update(status='В эксплуатации',target_kind='Сотрудник',target_id=1,owner='Первый',responsible_id=1,responsible='Первый',location='Офис')
            if id==4:a.update(status='В эксплуатации',target_kind='Объект',target_id='OBJ-001',owner='Кабинет',responsible_id=1,responsible='Первый',location='Офис')
            assets.append([a[k] for k in FIELDS])
        seed={'equipment':assets,'employees':[[1,'Первый','Инженер','Действующий'],[2,'Второй','Техник','Действующий'],[3,'Уволенный','Техник','Уволен']], 'objects':[['OBJ-001','Кабинет','Объект']], 'types':['Ноутбук'],'locations':['Склад','Офис','Этаж 2'],'issues':[],'legacy':[]}
        (res/'seed.json').write_text(json.dumps(seed,ensure_ascii=False),encoding='utf8')
        self.s=Store(root/'data',res);self.call('settings',operator='Оператор',role='Специалист')

    def call(self,method,**p):return getattr(self.s,method)({**p,'revision':self.s.snapshot()['revision']})
    def asset(self,id=3):return next(a for a in self.s.snapshot()['assets'] if a['id']==id)
    def transfer(self,**changes):
        p=dict(operation='Передача',ids=[3],target_kind='Сотрудник',target_id=2,responsible_id=2,location='Этаж 2',reason='Передача');p.update(changes);return self.call('movement',**p)
    def atomic_reject(self,method,**p):
        before=self.s.snapshot();files=set((self.s.root/'Акты').glob('*.docx'))
        with self.assertRaises(ValueError):self.call(method,**p)
        self.assertEqual(before,self.s.snapshot());self.assertEqual(files,set((self.s.root/'Акты').glob('*.docx')))

    def test_clear_responsible_correction(self):
        a=self.asset();self.call('correct_asset',id=3,reason='Снятие',responsible_id=None,**{k:a[k] for k in ['name','serial','kit','condition','notes']});self.assertFalse(self.asset()['responsible_id'])
    def test_transfer_employee(self):
        r=self.transfer();self.assertEqual(self.asset()['target_id'],2);self.assertEqual(self.asset()['responsible_id'],2);self.assertTrue((self.s.root/r['file']).exists())
    def test_transfer_object(self):
        self.transfer(target_kind='Объект',target_id='OBJ-001');self.assertEqual(self.asset()['target_kind'],'Объект')
    def test_transfer_same_owner_rejected(self):
        self.atomic_reject('movement',operation='Передача',ids=[3],target_kind='Сотрудник',target_id=1,responsible_id=1,location='Офис',reason='Тест')
    def test_transfer_dismissed_rejected(self):
        self.atomic_reject('movement',operation='Передача',ids=[3],target_kind='Сотрудник',target_id=3,responsible_id=1,location='Офис',reason='Тест')
    def test_transfer_stock_rejected(self):
        self.atomic_reject('movement',operation='Передача',ids=[1],target_kind='Сотрудник',target_id=2,responsible_id=2,location='Офис',reason='Тест')
    def test_move_preserves_assignment(self):
        a=self.asset();self.call('movement',operation='Перемещение',ids=[3],location='Этаж 2',reason='Переезд')
        for k in ['owner','target_id','target_kind','responsible_id','status']:self.assertEqual(a[k],self.asset()[k])
    def test_mark_broken_assigned(self):
        self.call('asset_status',id=3,status='Поломка',reason='Не включается');self.assertEqual(self.asset()['target_id'],1);self.assertEqual(self.asset()['status'],'Поломка')
    def test_repair_success_preserves_assignment(self):
        self.call('repair',id=3,phase='start',provider='Сервис',reason='Поломка');a=self.asset();self.assertEqual(a['status'],'В ремонте');self.assertEqual(a['target_id'],1)
        self.call('repair',id=3,phase='finish',result='Исправно',location='Этаж 2',condition='Исправное',reason='Заменен кабель');self.assertEqual(self.asset()['status'],'В эксплуатации');self.assertEqual(self.s.snapshot()['repairs'][0]['state'],'Завершен')
    def test_repair_stock_unsuccessful(self):
        self.call('repair',id=1,phase='start',provider='Сервис',reason='Поломка');self.call('repair',id=1,phase='finish',result='Неисправно',location='Склад',condition='Неисправное',reason='Не удалось');self.assertEqual(self.asset(1)['status'],'Поломка')
    def test_repair_stock_success(self):
        self.call('repair',id=1,phase='start',provider='Сервис',reason='Поломка');self.call('repair',id=1,phase='finish',result='Исправно',location='Склад',condition='Исправное',reason='Готово');self.assertEqual(self.asset(1)['status'],'На складе')
    def test_repair_duplicate_rejected(self):
        self.call('repair',id=3,phase='start',provider='Сервис',reason='Поломка');self.atomic_reject('repair',id=3,phase='start',provider='Другой',reason='Тест')
    def test_transfer_in_repair_rejected(self):
        self.call('repair',id=3,phase='start',provider='Сервис',reason='Поломка');self.atomic_reject('movement',operation='Перемещение',ids=[3],location='Этаж 2',reason='Тест')
    def test_failed_write_rolls_back_movement_and_act(self):
        before=self.s.snapshot();files=set((self.s.root/'Акты').glob('*.docx'));old=self.s._put
        def fail(*args):raise RuntimeError('Injected')
        self.s._put=fail
        try:
            with self.assertRaises(RuntimeError):self.transfer()
        finally:self.s._put=old
        self.assertEqual(before,self.s.snapshot());self.assertEqual(files,set((self.s.root/'Акты').glob('*.docx')))
    def test_failed_write_rolls_back_repair(self):
        before=self.s.snapshot();old=self.s._put
        def fail(*args):raise RuntimeError('Injected')
        self.s._put=fail
        try:
            with self.assertRaises(RuntimeError):self.call('repair',id=3,phase='start',provider='Сервис',reason='Тест')
        finally:self.s._put=old
        self.assertEqual(before,self.s.snapshot());self.assertFalse(list((self.s.root/'Акты').glob('*.docx')))
    def test_quality_detects_corrupt_records(self):
        with contextlib.closing(self.s._connect()) as c:
            a=self.s._get(c,'assets',3);a.update(target_id=999,serial='SN-1');self.s._put(c,'assets',a)
            b=self.s._get(c,'assets',2);b.update(target_id=1,owner='Первый');self.s._put(c,'assets',b);c.commit()
        messages=[x['message'] for x in self.s.snapshot()['quality']['assets']]
        for fragment in ['Повтор','отсутствует','На складе']:self.assertTrue(any(fragment in m for m in messages))
    def test_full_backup_and_restore(self):
        self.transfer();shutil.copyfile(self.s.resources/'Шаблон_акта.docx',self.s.root/'Шаблон_акта.docx')
        secret=self.s.root/'Подключение_Битрикс.json';secret.write_text('DO NOT EXPORT',encoding='utf8')
        r=self.s.backup();plan=self.s.restore_preview(r['file'])
        with zipfile.ZipFile(r['file']) as z:self.assertIn('Шаблон_акта.docx',z.namelist());self.assertNotIn(secret.name,z.namelist());self.assertEqual(plan['documents'],1)
        self.call('movement',operation='Перемещение',ids=[3],location='Офис',reason='После копии');oldrev=self.s.snapshot()['revision']
        result=self.s.restore({**plan,'revision':oldrev});self.assertEqual(self.asset()['location'],'Этаж 2');self.assertTrue(Path(result['safety_backup']).exists());self.assertGreater(self.s.snapshot()['revision'],oldrev);self.assertEqual(secret.read_text(),'DO NOT EXPORT')
    def test_stale_restore_rejected(self):
        plan=self.s.restore_preview(self.s.backup()['file']);old=self.s.snapshot()['revision'];self.call('employee',name='Новый',role='Инженер')
        with self.assertRaises(ValueError):self.s.restore({**plan,'revision':old})
    def test_archive_traversal_rejected(self):
        f=self.s.root/'evil.zip';files={'../evil.docx':hashlib.sha256(b'evil').hexdigest(),'Учет.sqlite':'bad'}
        with zipfile.ZipFile(f,'w') as z:
            z.writestr('manifest.json',json.dumps({'format':'matuchet-backup-v1','files':files}));z.writestr('../evil.docx',b'evil');z.writestr('Учет.sqlite',b'bad')
        with self.assertRaises(ValueError):self.s.restore_preview(f)
    def test_bad_template_does_not_commit(self):
        d=Document();d.add_paragraph('No table');d.save(self.s.root/'Шаблон_акта.docx');before=self.s.snapshot()
        with self.assertRaises(ValueError):self.transfer()
        self.assertEqual(before,self.s.snapshot())
    def test_table_selection_ignores_intro_table(self):
        d=Document(self.s.resources/'Шаблон_акта.docx');table=d.add_table(rows=1,cols=1);table.cell(0,0).text='Служебная таблица';d._element.body.insert(0,table._tbl);d.save(self.s.root/'Шаблон_акта.docx')
        r=self.transfer();out=Document(self.s.root/r['file']);self.assertEqual(out.tables[0].cell(0,0).text,'Служебная таблица');self.assertEqual(len(out.tables[1].rows),2)
    def test_remove_org_everywhere(self):
        d=Document();d.add_paragraph('Организация: Тест');d.add_table(rows=1,cols=1).cell(0,0).text='{{ORG}}';d.sections[0].header.paragraphs[0].text='Организация: {{ORG}}';without_organization(d)
        self.assertFalse(any('ORG' in p.text or 'Организация' in p.text for p in paragraphs(d)))
    def test_marker_preserves_runs(self):
        d=Document();p=d.add_paragraph();p.add_run('Начало ').bold=True;p.add_run('{{OPER');p.add_run('ATOR}}');p.add_run(' конец').italic=True;replace_text(p,'{{OPERATOR}}','Образец');self.assertEqual(p.text,'Начало Образец конец');self.assertTrue(p.runs[0].bold);self.assertTrue(p.runs[-1].italic)
    def test_unknown_marker_rejected(self):
        d=Document(self.s.resources/'Шаблон_акта.docx');d.sections[0].footer.paragraphs[0].text='{{UNKNOWN}}';d.save(self.s.root/'Шаблон_акта.docx')
        with self.assertRaises(ValueError):self.transfer()
    def test_excel_roundtrip_repair(self):
        self.call('repair',id=3,phase='start',provider='Сервис',reason='Тест');self.assertEqual(self.s.import_preview(self.s.export()['file'])['changes'],[])
    def test_old_database_no_repairs_meta(self):self.assertEqual(self.s.snapshot()['repairs'],[])
    def test_packaged_seed_consistency(self):
        seed=json.loads((RESOURCES/'seed.json').read_text(encoding='utf8'));s=Store(self.s.root.parent/'package_data',RESOURCES);d=s.snapshot()
        self.assertEqual(len(d['assets']),len(seed['equipment']));self.assertEqual(len(d['employees']),len(seed['employees']));self.assertEqual(len(d['legacy']),len(seed['legacy']))
    def test_restore_changed_archive_rejected(self):
        plan=self.s.restore_preview(self.s.backup()['file'])
        with zipfile.ZipFile(plan['file'],'a') as z:z.comment=b'Changed since preview'
        before=self.s.snapshot()
        with self.assertRaises(ValueError):self.s.restore({**plan,'revision':before['revision']})
        self.assertEqual(before,self.s.snapshot())
    def test_stale_mutation(self):
        with self.assertRaises(ValueError):self.s.movement({'revision':0,'operation':'Перемещение','ids':[3],'location':'Этаж 2','reason':'Тест'})
    def test_legacy_issue_return_and_dismissal(self):
        self.call('operation',operation='Выдача',target_kind='Сотрудник',target_id=2,responsible_id=2,location='Офис',items=[{'id':1,'action':'Выдача','kit':'Кабель','condition':'Б/у'}]);self.assertEqual(self.asset(1)['target_id'],2)
        self.call('operation',operation='Увольнение',target_kind='Сотрудник',target_id=2,items=[{'id':1,'action':'Возврат','kit':'Кабель','condition':'Б/у','status':'На складе'}]);self.assertEqual(self.asset(1)['status'],'На складе')
    @unittest.skipUnless(sys.platform=='win32','Windows DPAPI')
    def test_bitrix_secret_encrypted_and_upload_payload(self):
        b=BitrixDisk(self.s);webhook='https://example.org/rest/1/test-secret/'
        b.save({'webhook':webhook,'folder_id':123,'auto':False});self.assertNotIn('test-secret',b.path.read_text());self.assertNotIn('secret',b.public());self.assertEqual(protect(b.config()['secret'],True),webhook)
        calls=[]
        def fake(method,p):calls.append((method,p));return {'ID':'456'}
        b.call=fake;r=b.upload();self.assertEqual(r['file_id'],'456');self.assertEqual(calls[0][0],'disk.folder.uploadFile');self.assertTrue(calls[0][1]['generateUniqueName']);self.assertEqual(calls[0][1]['id'],123)
    def test_bitrix_failure_does_not_change_database(self):
        b=BitrixDisk(self.s);before=self.s.snapshot()
        with self.assertRaises(ValueError):b.upload()
        self.assertEqual(before,self.s.snapshot())
    @unittest.skipUnless(sys.platform=='win32','Windows DPAPI')
    def test_bitrix_download_foreign_host_rejected(self):
        b=BitrixDisk(self.s);b.save({'webhook':'https://example.org/rest/1/test-secret/','folder_id':123})
        b.call=lambda method,p:{'PARENT_ID':'123','NAME':'Матучет_test.zip','DOWNLOAD_URL':'https://other.example/test.zip'}
        with self.assertRaises(ValueError):b.download(456)
    @unittest.skipUnless(sys.platform=='win32','Windows DPAPI')
    def test_auto_upload_flush_on_close(self):
        b=BitrixDisk(self.s);b.save({'webhook':'https://example.org/rest/1/test-secret/','folder_id':123,'auto':True});calls=[];b.upload=lambda:calls.append('uploaded')
        b.schedule();b.close(flush=True);self.assertEqual(calls,['uploaded']);self.assertIsNone(b.timer)

def run(resources,root):
    global RESOURCES,TEST_ROOT
    RESOURCES=Path(resources);TEST_ROOT=Path(root)/'regression';TEST_ROOT.mkdir(parents=True,exist_ok=True)
    stream=io.StringIO();result=unittest.TextTestRunner(stream=stream,verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(ProjectTests))
    report={'run':result.testsRun,'success':result.wasSuccessful(),'failures':len(result.failures),'errors':len(result.errors),'details':stream.getvalue()}
    (TEST_ROOT.parent/'regression.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8');print(stream.getvalue());return 0 if result.wasSuccessful() else 1

if __name__=='__main__':sys.exit(run(Path(__file__).parent,Path(__file__).parent/'test'))
