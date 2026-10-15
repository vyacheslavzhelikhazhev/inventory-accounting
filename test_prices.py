import unittest, tempfile, json, shutil, hashlib, copy, io
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from project_tests import ProjectTests
from prices import price, money, total
from handover import fill_handover, AUTHORITY

ROOT=Path(__file__).parent

class PriceTests(ProjectTests):
    def setUp(self):
        super().setUp()
        shutil.copyfile(ROOT/'Шаблон_приема_передачи.docx',self.s.resources/'Шаблон_приема_передачи.docx')

    def issue(self,ids=(1,),target_kind='Сотрудник',target_id=2):
        return self.call('operation',operation='Выдача',target_kind=target_kind,target_id=target_id,responsible_id=2,location='Офис',items=[dict(id=i,action='Выдача',kit='Полный комплект',condition='Исправно') for i in ids])

    def set_price(self,id,value):
        a=self.asset(id)
        return self.call('correct_asset',id=id,reason='Цена по счету',purchase_price=value,**{k:a[k] for k in ['name','serial','kit','condition','notes']})

    def act(self,result):return Document(self.s.root/result['file'])

    def test_price_empty_zero_and_precision(self):
        self.assertEqual(price(''),'');self.assertEqual(price(None),'');self.assertEqual(price('0'),'0.00')
        self.assertEqual(price('11 799,50'),'11799.50');self.assertEqual(money('11799.50'),'11 799,50')
        for value in ['-1','1.001','NaN','Infinity','1e3','123 рублей','1000000000',True]:
            with self.assertRaises(ValueError):price(value)

    def test_empty_prices_and_blank_total(self):
        d=self.act(self.issue())
        self.assertEqual(d.tables[0].rows[1].cells[4].text,'')
        text='\n'.join(p.text for p in d.paragraphs)
        self.assertIn('на сумму ________________ руб.',text);self.assertNotIn('0,00',text)
        self.assertNotIn('коммерческого директора',text);self.assertIn('представитель:',text)
        self.assertNotIn('Оператор',text);self.assertIn('________________________________',text)

    def test_employer_settings_fill_new_act_and_preserve_old_act(self):
        self.call('settings',operator='Оператор',role='Специалист',employer_representative='Образцов Пример Примерович',employer_role='Руководитель',authority='Тестовый документ 123')
        r=self.issue();old=(self.s.root/r['file']).read_bytes()
        text='\n'.join(p.text for p in self.act(r).paragraphs)
        for value in ['Образцов Пример Примерович','Руководитель','Тестовый документ 123','Образцов П.П.']:
            self.assertIn(value,text)
        self.call('settings',operator='Другой',role='Специалист')
        self.assertEqual(self.s.snapshot()['settings']['employer_representative'],'Образцов Пример Примерович')
        self.assertEqual((self.s.root/r['file']).read_bytes(),old)

    def test_public_seed_is_empty_and_new_database_stays_empty(self):
        from store import Store
        seed=json.loads((ROOT/'seed.json').read_text(encoding='utf-8'))
        for key in ['equipment','employees','objects','issues','legacy']:
            self.assertEqual(seed[key],[])
        s=Store(self.s.root.parent/'empty-public',ROOT).snapshot()
        for key in ['assets','employees','acts','history']:
            self.assertEqual(s[key],[])
        self.assertFalse(s['settings'].get('employer_representative'))

    def test_known_and_partial_total(self):
        self.set_price(1,'11799,00')
        d=self.act(self.issue((1,2)))
        self.assertEqual(d.tables[0].rows[1].cells[4].text,'11 799,00')
        self.assertEqual(d.tables[0].rows[2].cells[4].text,'')
        self.assertTrue(any('на сумму ________________ руб.' in p.text for p in d.paragraphs))

    def test_known_total_decimal_and_old_act_immutable(self):
        self.set_price(1,'0.10');self.set_price(2,'0.20');r=self.issue((1,2));path=self.s.root/r['file'];old=path.read_bytes()
        self.assertTrue(any('на сумму 0,30 руб.' in p.text for p in Document(path).paragraphs))
        self.set_price(1,'999');self.assertEqual(old,path.read_bytes())

    def test_return_uses_reverse_wording_and_stored_price(self):
        self.set_price(3,'500');r=self.call('operation',operation='Возврат',target_kind='Сотрудник',target_id=1,items=[dict(id=3,action='Возврат',kit='Кабель',condition='Б/у')])
        d=self.act(r);self.assertTrue(any('Работник передал, а Работодатель принял на склад' in p.text for p in d.paragraphs));self.assertEqual(d.tables[0].rows[1].cells[4].text,'500,00')

    def test_shared_object_has_responsible_signatory(self):
        d=self.act(self.issue(target_kind='Объект',target_id='OBJ-001'))
        text='\n'.join(p.text for p in d.paragraphs)
        self.assertIn('и Второй, проживающий',text);self.assertIn('Место общего использования: Кабинет.',text)

    def test_replacement_separate_totals(self):
        self.set_price(3,'100');self.set_price(1,'200')
        d=self.act(self.call('operation',operation='Замена',target_kind='Сотрудник',target_id=1,responsible_id=1,location='Офис',items=[dict(id=3,action='Возврат',kit='Кабель',condition='Б/у'),dict(id=1,action='Выдача',kit='Кабель',condition='Новое')]))
        text='\n'.join(p.text for p in d.paragraphs)
        self.assertIn('Возвращено: 1 шт., на сумму 100,00 руб.',text);self.assertIn('Выдано взамен: 1 шт., на сумму 200,00 руб.',text)

    def test_receive_and_export_import_prices(self):
        self.call('receive',mode='Новая поставка',document='Счет 123',items=[dict(type='Ноутбук',name='Новый',serial='PRICE-SN',kit='Кабель',condition='Новое',purchase_price='11799,50')])
        a=self.s.snapshot()['assets'][-1];self.assertEqual(a['purchase_price'],'11799.50')
        f=self.s.export()['file'];w=load_workbook(f);ws=w['Оборудование'];col=[c.value for c in ws[1]].index('purchase_price')+1;ws.cell(2,col,123.45);w.save(f);w.close()
        plan=self.s.import_preview(f);self.call('import_commit',file=f,preview=plan,reason='Цены по документам');self.assertEqual(self.asset(1)['purchase_price'],'123.45')
        # Earlier exports lack the new column: importing them must not clear a known price.
        f=self.s.export()['file'];w=load_workbook(f);w['Оборудование'].delete_cols(col);w.save(f);w.close();self.assertEqual(self.s.import_preview(f)['changes'],[])

    def test_invalid_price_does_not_mutate(self):
        a=self.asset(1);self.atomic_reject('correct_asset',id=1,reason='Тест',purchase_price='-10',**{k:a[k] for k in ['name','serial','kit','condition','notes']})
        self.atomic_reject('receive',mode='Новая поставка',document='x',items=[dict(type='Ноутбук',name='x',serial='new',kit='x',condition='x',purchase_price='1.234')])

    def test_new_external_template_is_backed_up_and_restored(self):
        template=self.s.root/'Шаблон_приема_передачи.docx';shutil.copyfile(self.s.resources/template.name,template)
        orig=template.read_bytes();p=self.s.restore_preview(self.s.backup()['file']);template.write_bytes(b'changed')
        self.s.restore({**p,'revision':self.s.snapshot()['revision']});self.assertEqual(template.read_bytes(),orig)

    def test_bad_handover_rolls_back(self):
        d=Document();d.add_paragraph('Invalid');d.save(self.s.root/'Шаблон_приема_передачи.docx');before=self.s.snapshot()
        with self.assertRaises(ValueError):self.issue()
        self.assertEqual(before,self.s.snapshot());self.assertEqual(list((self.s.root/'Акты').glob('*.docx')),[])

def run(resources,test_root):
    global ROOT
    import project_tests
    ROOT=Path(resources);project_tests.RESOURCES=ROOT
    project_tests.TEST_ROOT=Path(test_root)/'regression';project_tests.TEST_ROOT.mkdir(parents=True,exist_ok=True)
    suite=unittest.defaultTestLoader.loadTestsFromTestCase(PriceTests)
    stream=io.StringIO();result=unittest.TextTestRunner(stream=stream,verbosity=2).run(suite)
    (Path(test_root)/'regression.json').write_text(json.dumps({'run':result.testsRun,'success':result.wasSuccessful(),'details':stream.getvalue()},ensure_ascii=False,indent=2),encoding='utf8')
    print(stream.getvalue());return 0 if result.wasSuccessful() else 1

if __name__=='__main__':raise SystemExit(run(ROOT,ROOT/'test'))
