import copy
import datetime as dt
import json
import os
from pathlib import Path
import sqlite3
import threading
import uuid
from prices import price
from handover import fill_handover
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FIELDS = ['id','type','name','serial','status','kit','notes','responsible','location','owner','purpose','keyboard','condition','purchase_date','supply_doc','target_kind','target_id','responsible_id','updated','warnings','purchase_price']
def now(): return dt.datetime.now().isoformat(timespec='seconds')
def clean(v): return str(v or '').strip()
def require(ok, message):
    if not ok: raise ValueError(message)

from enhancements import without_organization, paragraphs as all_paragraphs, replace_text, equipment_table, validate_markers
import hashlib
import zipfile
import tempfile
import shutil
import contextlib

class ClosingConnection(sqlite3.Connection):
    def __exit__(self,*args):
        try:return super().__exit__(*args)
        finally:self.close()

class Store:
    def __init__(self, root, resources):
        self.root, self.resources = Path(root), Path(resources)
        self.root.mkdir(parents=True, exist_ok=True)
        for folder in ['Акты','Резервные копии','Выгрузки']: (self.root/folder).mkdir(exist_ok=True)
        self.db = self.root/'Учет.sqlite'
        self.lock = threading.RLock()
        self._initialize()

    def _connect(self):
        con = sqlite3.connect(self.db, timeout=10, factory=ClosingConnection)
        con.row_factory = sqlite3.Row
        con.execute('PRAGMA foreign_keys=ON')
        con.execute('PRAGMA journal_mode=DELETE')
        return con

    def _initialize(self):
        with self._connect() as c:
            c.executescript('''CREATE TABLE IF NOT EXISTS employees(id INTEGER PRIMARY KEY,data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS assets(id INTEGER PRIMARY KEY,data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY,data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS history(id INTEGER PRIMARY KEY AUTOINCREMENT, stamp TEXT, operation TEXT, target TEXT, actor TEXT, items TEXT, detail TEXT);
            CREATE TABLE IF NOT EXISTS acts(id INTEGER PRIMARY KEY AUTOINCREMENT, stamp TEXT, operation TEXT, target TEXT, file TEXT, status TEXT);''')
            if not c.execute("SELECT 1 FROM meta WHERE key='seeded'").fetchone():
                m = json.loads((self.resources/'seed.json').read_text(encoding='utf-8'))
                for r in m['employees']:
                    self._put(c,'employees',{'id':r[0],'name':r[1],'role':r[2],'status':r[3]})
                for r in m['equipment']: self._put(c,'assets',dict(zip(FIELDS,r)))
                for key in ['objects','issues','types','locations','legacy']: self._meta(c,key,m[key])
                self._meta(c,'settings',{'operator':'','role':''})
                self._meta(c,'seeded',True)
                self._meta(c,'revision',0)
                max_id=max([0]+[r[0] for r in m['equipment']]+[int(r[2]) for r in m['legacy'] if str(r[2]).isdigit()])
                self._meta(c,'next_asset',max_id+1)

    def _meta(self,c,key,value=None):
        if value is not None: c.execute('INSERT OR REPLACE INTO meta VALUES (?,?)',(key,json.dumps(value,ensure_ascii=False))); return value
        row=c.execute('SELECT data FROM meta WHERE key=?',(key,)).fetchone()
        return json.loads(row[0]) if row else None

    def _put(self,c,table,item): c.execute(f'INSERT OR REPLACE INTO {table} VALUES (?,?)',(item['id'],json.dumps(item,ensure_ascii=False)))
    def _all(self,c,table): return [json.loads(r['data']) for r in c.execute(f'SELECT data FROM {table} ORDER BY id')]
    def _get(self,c,table,id):
        row=c.execute(f'SELECT data FROM {table} WHERE id=?',(id,)).fetchone()
        require(row is not None,'Запись больше не существует. Обновите список.')
        return json.loads(row[0])

    def snapshot(self):
        with self.lock, self._connect() as c:
            return {'employees':self._all(c,'employees'),'assets':self._all(c,'assets'),
                    'acts':[dict(r) for r in c.execute('SELECT * FROM acts ORDER BY id DESC')],
                    'history':[dict(r) for r in c.execute('SELECT * FROM history ORDER BY id DESC')],
                    **{k:self._meta(c,k) for k in ['objects','types','locations','issues','settings','revision','legacy']}, 'quality':self._quality(c), 'repairs':self._meta(c,'repairs') or [], 'version':'1.5.1'}

    def _backup(self,c):
        # All callers run this before beginning their write transaction.
        return Path(self.backup()['file'])

    def backup(self):
        with self.lock, contextlib.closing(self._connect()) as c:
            stamp=dt.datetime.now().strftime('%Y-%m-%d_%H-%M-%S')+'_'+uuid.uuid4().hex[:8]
            file=self.root/'Резервные копии'/('Матучет_'+stamp+'.zip')
            with tempfile.TemporaryDirectory(prefix='matuchet-backup-',dir=self.root) as folder:
                db=Path(folder)/'Учет.sqlite'
                with contextlib.closing(sqlite3.connect(db)) as dest:c.backup(dest)
                files={'Учет.sqlite':db}
                for path in (self.root/'Акты').rglob('*.docx'):
                    require(not path.is_symlink(),'Ссылка в папке актов не поддерживается.')
                    require(path.resolve().is_relative_to((self.root/'Акты').resolve()),'Файл акта находится вне папки актов.')
                    files[path.relative_to(self.root).as_posix()]=path
                for name in ['Шаблон_акта.docx','Шаблон_приема_передачи.docx']:
                    template=self.root/name
                    if template.exists():files[name]=template
                manifest={'format':'matuchet-backup-v1','version':'1.5.1','created':now(),'revision':self._meta(c,'revision'),'files':{name:hashlib.sha256(path.read_bytes()).hexdigest() for name,path in files.items()}}
                pending=file.with_suffix('.pending')
                try:
                    with zipfile.ZipFile(pending,'w',zipfile.ZIP_DEFLATED) as z:
                        for name,path in files.items():z.write(path,name)
                        z.writestr('manifest.json',json.dumps(manifest,ensure_ascii=False))
                    self.restore_preview(pending)
                    os.replace(pending,file)
                finally:pending.unlink(missing_ok=True)
            return {'file':str(file),'revision':manifest['revision']}

    def _revision(self,c,p):
        require(p.get('revision')==self._meta(c,'revision'),'Данные изменились. Обновите окно и повторите действие.')
    def _finish(self,c): self._meta(c,'revision',self._meta(c,'revision')+1)
    def _log(self,c,operation,target,items,detail=''):
        s=self._meta(c,'settings')
        actor=s['operator'] or os.environ.get('USERNAME','')
        c.execute('INSERT INTO history(stamp,operation,target,actor,items,detail) VALUES (?,?,?,?,?,?)',
                  (now(),operation,target,actor,json.dumps(items,ensure_ascii=False),detail))

    def settings(self,p):
        with self.lock, self._connect() as c:
            self._revision(c,p); self._backup(c)
            self._meta(c,'settings',{k:clean(p.get(k)) for k in ['operator','role'] } | {k:clean(p.get(k,self._meta(c,'settings').get(k,''))) for k in ['authority','employer_representative','employer_role']})
            self._log(c,'Настройки','',[], 'Обновлены реквизиты актов'); self._finish(c)
        return {}

    def employee(self,p):
        with self.lock, self._connect() as c:
            self._revision(c,p)
            name,role=clean(p.get('name')),clean(p.get('role'))
            require(name and role,'Заполните имя и должность.')
            people=self._all(c,'employees')
            require(not any(e['name'].casefold()==name.casefold() and e['id']!=p.get('id') for e in people),'Сотрудник с таким именем уже существует. Уточните полное ФИО.')
            self._backup(c)
            if p.get('id'):
                e=self._get(c,'employees',p['id']); old=copy.deepcopy(e)
                e.update(name=name,role=role)
                self._put(c,'employees',e)
                for a in self._all(c,'assets'):
                    changed=False
                    if a['target_kind']=='Сотрудник' and a['target_id']==e['id']: a['owner']=name;changed=True
                    if a['responsible_id']==e['id']: a['responsible']=name;changed=True
                    if changed: self._put(c,'assets',a)
                self._log(c,'Изменение сотрудника',name,[{'before':old,'after':e}])
            else:
                e={'id':max([e['id'] for e in people]+[0])+1,'name':name,'role':role,'status':'Действующий'}
                self._put(c,'employees',e); self._log(c,'Новый сотрудник',name,[{'after':e}])
            self._finish(c)
        return {'id':e['id']}

    def _recipient(self,c,p):
        kind=p.get('target_kind','Сотрудник'); id=p.get('target_id')
        if kind=='Сотрудник':
            e=self._get(c,'employees',id)
            return kind,id,e['name'],e
        require(kind=='Объект','Выберите сотрудника или объект.')
        obj=next((o for o in self._meta(c,'objects') if o[0]==id),None)
        require(obj is not None,'Выберите объект.')
        return kind,id,obj[1],None

    def _act(self,c,operation,target,items,p):
        s=self._meta(c,'settings')
        require(all(clean(s.get(k)) for k in ['operator','role']),'Заполните ФИО и должность выдающего в настройках.')
        number=(c.execute('SELECT COALESCE(MAX(id),0)+1 FROM acts').fetchone()[0])
        filename=f'{number:05d}_{dt.date.today().isoformat()}_{operation}.docx'
        relative='Акты/'+filename
        path=self.root/relative
        require(not path.exists(),'Файл акта с таким номером уже существует. Проверьте папку актов.')
        handover_name='Шаблон_приема_передачи.docx'
        handover_template=self.root/handover_name
        if not handover_template.exists():handover_template=self.resources/handover_name
        if operation in ['Выдача','Возврат','Увольнение','Замена'] and handover_template.exists():
            d=Document(handover_template)
            fill_handover(d,str(number).zfill(5),dt.date.today().strftime('%d.%m.%Y'),operation,target,items,p,s)
        else:
            template=self.root/'Шаблон_акта.docx'
            if not template.exists(): template=self.resources/'Шаблон_акта.docx'
            require(template.exists(),'Не найден шаблон акта.')
            d=Document(template)
            without_organization(d)
            values={'NUMBER':str(number).zfill(5),'DATE':dt.date.today().strftime('%d.%m.%Y'),
                    'OPERATION':operation,'EMPLOYEE':target,'OPERATOR':s['operator'],'ROLE':s['role'],
                    'BASIS':clean(p.get('basis')) or '—','NOTE':clean(p.get('note')) or '—'}
            paragraphs=list(all_paragraphs(d))
            for par in paragraphs:
                for token,value in values.items():
                    marker='{{'+token+'}}'
                    if marker in par.text:
                        replace_text(par,marker,value)
            validate_markers(d)
            table=equipment_table(d)
            while len(table.rows)>1: table._tbl.remove(table.rows[-1]._tr)
            for entry in items:
                a=entry['after']; action=entry['action']
                row=table.add_row()
                values=[str(a['id']),action,a['type']+'\n'+a['name'],clean(a['serial']) or 'Не указан',clean(a['kit']) or 'Не уточнена',(clean(a['condition']) or 'Не уточнено')+'\n'+clean(entry.get('note'))]
                for cell,value in zip(row.cells,values):
                    cell.text=value
                    for par in cell.paragraphs:
                        for run in par.runs: run.font.size=Pt(9);run.bold=False
                trPr=row._tr.get_or_add_trPr(); noSplit=OxmlElement('w:cantSplit');trPr.append(noSplit)
        temp=path.with_suffix('.pending.docx')
        try: d.save(temp); os.replace(temp,path)
        except Exception:
            temp.unlink(missing_ok=True); raise
        try:
            c.execute('INSERT INTO acts(stamp,operation,target,file,status) VALUES (?,?,?,?,?)',(now(),operation,target,relative,'Ожидает подписи'))
        except Exception:
            path.unlink(missing_ok=True);raise
        return relative

    def operation(self,p):
        path=None
        with self.lock:
            c=self._connect()
            try:
                self._revision(c,p)
                op=p.get('operation'); require(op in ['Выдача','Возврат','Замена','Увольнение'],'Неизвестная операция.')
                kind,id,name,e=self._recipient(c,p)
                if op in ['Выдача','Замена','Увольнение']: require(e is None or e['status']=='Действующий','Сотрудник уволен.')
                require(op!='Увольнение' or e is not None,'Увольнение доступно только для сотрудника.')
                all_assets=self._all(c,'assets')
                responsibility=[]
                owned=[a for a in all_assets if a['target_kind']==kind and a['target_id']==id]
                rows=p.get('items',[])
                require(len(rows)<=100,'Не более 100 позиций в одной операции.')
                require(rows or op=='Увольнение','Выберите оборудование.')
                ids=[r.get('id') for r in rows]
                require(len(ids)==len(set(ids)),'Одна позиция добавлена дважды.')
                returns=[r for r in rows if r.get('action')=='Возврат']; issues=[r for r in rows if r.get('action')=='Выдача']
                require(len(returns)+len(issues)==len(rows),'У каждой позиции должно быть действие.')
                if op=='Замена': require(returns and issues,'Для замены выберите возвращаемую и выдаваемую технику.')
                if op=='Выдача': require(not returns,'В выдаче не должно быть возвратов.')
                if op in ['Возврат','Увольнение']: require(not issues,'В возврате не должно быть выдачи.')
                if op=='Увольнение':
                    require(set(ids)=={a['id'] for a in owned},'При увольнении нужно вернуть все закрепленное оборудование.')
                    require(not any(a['target_kind']=='Требует уточнения' and clean(a['owner']).casefold()==name.casefold() for a in all_assets),'Есть оборудование с неуточненным владельцем. Сначала уточните его в карточке оборудования.')
                    unresolved=[a for a in all_assets if not a['responsible_id'] and clean(a['responsible']).casefold()==name.casefold()]
                    if unresolved:
                        require(sum(clean(x['name']).casefold()==name.casefold() for x in self._all(c,'employees'))==1,'Есть оборудование с неоднозначным ответственным. Сначала уточните ответственного в карточке оборудования.')
                    responsibility=[a for a in all_assets if a['responsible_id']==id or a in unresolved]
                entries=[]
                resp=None
                if issues:
                    require(p.get('location') in self._meta(c,'locations'),'Выберите расположение.')
                    resp=self._get(c,'employees',p.get('responsible_id'))
                    require(resp['status']=='Действующий','Выберите действующего ответственного.')
                for r in rows:
                    a=self._get(c,'assets',r['id']); before=copy.deepcopy(a)
                    require(clean(r.get('kit')) and clean(r.get('condition')),'Укажите комплектность и состояние каждой позиции.')
                    if r['action']=='Выдача':
                        require(a['status']=='На складе' and not a['target_id'] and not clean(a['owner']),'Эта позиция уже закреплена или недоступна на складе.')
                        if clean(a['serial']): require(sum(clean(x['serial']).casefold()==clean(a['serial']).casefold() for x in all_assets)==1,'Повтор серийного номера. Сначала уточните данные.')
                        a.update(status='В эксплуатации',target_kind=kind,target_id=id,owner=name,responsible=resp['name'],responsible_id=resp['id'],location=p['location'])
                    else:
                        require(a['target_kind']==kind and a['target_id']==id,'Оборудование закреплено за другим получателем.')
                        require(a['status'] in ['В эксплуатации','Поломка'],'Эту позицию нельзя вернуть из текущего статуса.')
                        status=r.get('status','На складе');require(status in ['На складе','Поломка'],'Выберите статус возврата.')
                        a.update(status=status,target_kind='',target_id='',owner='',responsible='',responsible_id='',location='Склад')
                    a.update(kit=clean(r['kit']),condition=clean(r['condition']),updated=now())
                    entries.append({'action':r['action'],'before':before,'after':a,'note':clean(r.get('note'))})
                self._backup(c);c.execute('BEGIN IMMEDIATE')
                if entries: path=self._act(c,op,name,entries,p)
                for entry in entries: self._put(c,'assets',entry['after'])
                released=[]
                if op=='Увольнение':
                    # Returned assets were already cleared; other recipients retain their equipment.
                    returned_ids={entry['after']['id'] for entry in entries}
                    for a in responsibility:
                        if a['id'] in returned_ids:continue
                        before=copy.deepcopy(a)
                        a.update(responsible='',responsible_id='',updated=now())
                        self._put(c,'assets',a)
                        released.append({'action':'Снятие ответственности','before':before,'after':a})
                    e['status']='Уволен';self._put(c,'employees',e)
                    if released:self._log(c,'Снятие ответственности',name,released,'Увольнение сотрудника. Получатели, расположение и статус оборудования сохранены.')
                self._log(c,op,name,entries,clean(p.get('note')));self._finish(c);c.commit()
                return {'file':path,'count':len(entries),'released_responsibility':len(released)}
            except Exception:
                c.rollback()
                if path: (self.root/path).unlink(missing_ok=True)
                raise
            finally: c.close()

    def receive(self,p):
        path=None
        with self.lock:
            c=self._connect()
            try:
                self._revision(c,p); rows=p.get('items',[])
                require(rows and len(rows)<=100,'Добавьте от 1 до 100 позиций.')
                require(p.get('mode') in ['Новая поставка','Ранее не учтено'],'Выберите основание поступления.')
                if p['mode']=='Новая поставка': require(clean(p.get('document')),'Укажите номер документа поставки.')
                serials={clean(a['serial']).casefold() for a in self._all(c,'assets') if clean(a['serial'])}
                entries=[]; id=self._meta(c,'next_asset')
                for r in rows:
                    require(r.get('type') in self._meta(c,'types'),'Выберите тип оборудования.')
                    require(all(clean(r.get(k)) for k in ['name','kit','condition']),'Заполните название, комплектность и состояние.')
                    serial=clean(r.get('serial'))
                    if serial: require(serial.casefold() not in serials,'Серийный номер уже существует: '+serial);serials.add(serial.casefold())
                    a={k:'' for k in FIELDS}
                    a.update(purchase_price=price(r.get('purchase_price')),id=id,type=r['type'],name=clean(r['name']),serial=serial,status='На складе',location='Склад',kit=clean(r['kit']),condition=clean(r['condition']),supply_doc=clean(p.get('document')),purchase_date=dt.date.today().isoformat() if p['mode']=='Новая поставка' else '',updated=now(),warnings='' if serial else 'Не указан серийный номер')
                    entries.append({'action':'Поступление','before':None,'after':a});id+=1
                self._backup(c);c.execute('BEGIN IMMEDIATE')
                path=self._act(c,'Поступление','Склад',entries,{'basis':p['mode']+' '+clean(p.get('document'))})
                for entry in entries: self._put(c,'assets',entry['after'])
                self._meta(c,'next_asset',id);self._log(c,'Поступление','Склад',entries,p['mode']);self._finish(c);c.commit()
                return {'file':path,'count':len(entries)}
            except Exception:
                c.rollback()
                if path: (self.root/path).unlink(missing_ok=True)
                raise
            finally: c.close()

    def responsible(self,p):
        with self.lock,self._connect() as c:
            self._revision(c,p);a=self._get(c,'assets',p['id']);before=copy.deepcopy(a)
            require(a['status']!='Списано','Нельзя назначить ответственного за списанное оборудование.')
            require(clean(p.get('reason')),'Укажите основание смены ответственного.')
            if p.get('responsible_id'):
                e=self._get(c,'employees',p['responsible_id'])
                require(e['status']=='Действующий','Выберите действующего сотрудника.')
                a.update(responsible=e['name'],responsible_id=e['id'])
            else:a.update(responsible='',responsible_id='')
            a['updated']=now()
            self._backup(c);self._put(c,'assets',a)
            self._log(c,'Смена ответственного',a['name'],[{'before':before,'after':a}],clean(p['reason']));self._finish(c)
        return {}

    def correct_asset(self,p):
        with self.lock, self._connect() as c:
            self._revision(c,p);a=self._get(c,'assets',p['id']);old=copy.deepcopy(a)
            require(clean(p.get('reason')),'Укажите причину исправления.')
            for k in ['name','serial','kit','condition','notes']: a[k]=clean(p.get(k))
            if 'purchase_price' in p:a['purchase_price']=price(p['purchase_price'])
            require(a['name'] and a['kit'] and a['condition'],'Заполните название, комплектность и состояние.')
            if a['serial']: require(not any(x['id']!=a['id'] and clean(x['serial']).casefold()==a['serial'].casefold() for x in self._all(c,'assets')),'Серийный номер уже существует.')
            # Owner correction only resolves ambiguous imported records; normal movements use operations.
            if old['target_kind']=='Требует уточнения' or (old['status']=='В эксплуатации' and not old['target_id']):
                if p.get('target_id'):
                    kind,id,name,e=self._recipient(c,p)
                    require(e is None or e['status']=='Действующий','Сотрудник уволен.')
                    a.update(target_kind=kind,target_id=id,owner=name)
            if p.get('responsible_id'):
                e=self._get(c,'employees',p['responsible_id']);require(e['status']=='Действующий','Ответственный уволен.')
                require(a['status']!='Списано','Нельзя назначить ответственного за списанное оборудование.')
                a.update(responsible=e['name'],responsible_id=e['id'])
            elif 'responsible_id' in p:a.update(responsible='',responsible_id='')
            a['updated']=now();a['warnings']=''
            self._backup(c);self._put(c,'assets',a);self._log(c,'Исправление данных',a['name'],[{'before':old,'after':a}],clean(p['reason']));self._finish(c)
        return {}

    def asset_status(self,p):
        with self.lock,self._connect() as c:
            self._revision(c,p);a=self._get(c,'assets',p['id']);old=copy.deepcopy(a)
            require(p.get('status') in ['Поломка','Списано','На складе'],'Недопустимый статус.')
            require(clean(p.get('reason')),'Укажите причину изменения статуса.')
            require(a['status']!='В ремонте','Сначала завершите ремонт.')
            if p['status']!='Поломка':require(not a['target_id'] and not clean(a['owner']),'Сначала оформите возврат от текущего получателя.')
            require(a['status']!='Списано','Списанное оборудование нельзя вернуть в учет этой операцией.')
            if p['status']=='На складе':require(a['status']=='Поломка','На склад можно вернуть оборудование после ремонта.')
            a.update(status=p['status'],updated=now())
            self._backup(c);self._put(c,'assets',a);self._log(c,p['status'],a['name'],[{'before':old,'after':a}],clean(p['reason']));self._finish(c)
        return {}

    def signed(self,p):
        with self.lock,self._connect() as c:
            self._revision(c,p);require(c.execute('SELECT 1 FROM acts WHERE id=?',(p['id'],)).fetchone(),'Акт не найден.')
            self._backup(c);c.execute("UPDATE acts SET status='Подписан' WHERE id=?",(p['id'],));self._log(c,'Акт подписан',str(p['id']),[]);self._finish(c)
        return {}

    def dictionary(self,p):
        with self.lock,self._connect() as c:
            self._revision(c,p);require(p.get('key') in ['types','locations','objects'],'Неизвестный справочник.')
            value=clean(p.get('value'));require(value,'Введите название.')
            key=p['key'];items=self._meta(c,key)
            names=[r[1] for r in items] if key=='objects' else items
            require(value.casefold() not in [s.casefold() for s in names],'Название уже существует.')
            self._backup(c)
            if key=='objects': items.append([f'OBJ-{len(items)+1:03d}',value,'Объект / общая техника'])
            else: items.append(value)
            self._meta(c,key,items);self._log(c,'Справочник',value,[],key);self._finish(c)
        return {}

    def export(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        s=self.snapshot(); w=Workbook();w.remove(w.active)
        def sheet(name,headers,rows):
            ws=w.create_sheet(name);ws.append(headers)
            for row in rows: ws.append([("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for v in row])
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='203C72')
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=min(55,max(12,max(len(str(cell.value or '')) for cell in col)+2))
            return ws
        sheet('Оборудование',FIELDS,[[a.get(k) for k in FIELDS] for a in s['assets']])
        sheet('Сотрудники',['id','name','role','status'],[[e[k] for k in ['id','name','role','status']] for e in s['employees']])
        sheet('Объекты',['id','name','kind'],s['objects'])
        sheet('История',['Дата','Операция','Получатель','Автор','Подробности'],[[h[k] for k in ['stamp','operation','target','actor','detail']] for h in s['history']])
        sheet('Старый журнал',['№','Дата','ID оборудования','Тип','Название','Операция','Документ','Получатель','Автор'],s['legacy'])
        sheet('Акты',['№','Дата','Операция','Получатель','Файл','Статус'],[[a[k] for k in ['id','stamp','operation','target','file','status']] for a in s['acts']])
        file=self.root/'Выгрузки'/('Матучет_'+dt.datetime.now().strftime('%Y-%m-%d_%H-%M-%S')+'.xlsx');w.save(file)
        return {'file':str(file)}

    def import_preview(self,file):
        from openpyxl import load_workbook
        w=load_workbook(file,read_only=True,data_only=False)
        try:
            require('Оборудование' in w.sheetnames and 'Сотрудники' in w.sheetnames,'Выберите выгрузку из этого приложения.')
            def read(name,keys):
                rows=list(w[name].values);headers=list(rows[0]);require(headers==keys or (name=='Оборудование' and headers==keys[:-1]),'Колонки выгрузки изменены: '+name)
                result=[]
                for r in rows[1:]:
                    if not any(v is not None for v in r):continue
                    item=dict(zip(headers,r))
                    for k,v in item.items():
                        if isinstance(v,str) and v.startswith("'") and len(v)>1 and v[1] in '=+-@':item[k]=v[1:]
                        if isinstance(v,str) and v.startswith('='):raise ValueError('Формулы при импорте недопустимы.')
                    require(isinstance(item['id'],int),'ID должен быть целым числом.')
                    result.append(item)
                return result
            assets=read('Оборудование',FIELDS);employees=read('Сотрудники',['id','name','role','status'])
        finally:w.close()
        s=self.snapshot(); ea={e['id']:e for e in employees};aa={a['id']:a for a in assets}
        require(len(ea)==len(employees) and len(aa)==len(assets),'Повтор ID в выгрузке.')
        require(set(ea)=={e['id'] for e in s['employees']} and set(aa)=={a['id'] for a in s['assets']},'Импорт поддерживает исправление существующих записей. Добавление и удаление оформляйте в приложении.')
        changes=[]
        for e in s['employees']:
            n=ea[e['id']]
            require(n['status']==e['status'],'Статус сотрудника меняется через операции.')
            if clean(n['name']).casefold()!=clean(e['name']).casefold():
                require(not any(other['id']!=n['id'] and clean(other['name']).casefold()==clean(n['name']).casefold() for other in employees),'Имя сотрудника уже существует. Уточните полное ФИО.')
            for k in ['name','role']:
                if clean(n[k])!=clean(e[k]):changes.append({'entity':'employees','id':e['id'],'label':e['name'],'field':k,'before':e[k],'after':clean(n[k])})
                require(clean(n[k]),'Имя и должность обязательны.')
        serials=[]
        for a in s['assets']:
            n=aa[a['id']]
            # Excel is for metadata exchange, never bypass the movement journal.
            for k in FIELDS:
                if k=='purchase_price':
                    if k in n and price(n[k])!=price(a.get(k)):
                        changes.append({'entity':'assets','id':a['id'],'label':a['name'],'field':k,'before':a.get(k,''),'after':price(n[k])})
                elif k in ['name','serial','kit','condition','notes']:
                    if clean(n[k])!=clean(a[k]):changes.append({'entity':'assets','id':a['id'],'label':a['name'],'field':k,'before':a[k],'after':clean(n[k])})
                else:require(clean(n[k])==clean(a[k]),'При импорте нельзя менять статус, привязки и служебные поля: '+a['name'])
            require(clean(n['name']),'Название обязательно.')
            for k in ['kit','condition']:
                if clean(n[k])!=clean(a[k]): require(clean(n[k]),'Комплектность и состояние нельзя очищать.')
            if clean(n['serial']):serials.append(clean(n['serial']).casefold())
        require(len(serials)==len(set(serials)),'Повтор серийного номера.')
        return {'revision':s['revision'],'changes':changes}

    def import_commit(self,p):
        # Re-read the file and compare preview to prevent changed-file imports.
        preview=self.import_preview(p['file']); require(preview==p['preview'],'Выгрузка или база изменились. Проверьте импорт заново.')
        require(clean(p.get('reason')),'Укажите причину импорта.')
        with self.lock,self._connect() as c:
            self._revision(c,preview);self._backup(c)
            entries=[]
            grouped={}
            for x in preview['changes']:grouped.setdefault((x['entity'],x['id']),[]).append(x)
            for (table,id),changes in grouped.items():
                a=self._get(c,table,id);before=copy.deepcopy(a)
                for x in changes:a[x['field']]=x['after']
                if table=='assets':a['updated']=now()
                self._put(c,table,a);entries.append({'entity':table,'before':before,'after':a})
            for a in self._all(c,'assets'):
                changed=False
                for entry in entries:
                    if entry['entity']=='employees':
                        e=entry['after']
                        if a['target_kind']=='Сотрудник' and a['target_id']==e['id']:a['owner']=e['name'];changed=True
                        if a['responsible_id']==e['id']:a['responsible']=e['name'];changed=True
                if changed:self._put(c,'assets',a)
            self._log(c,'Импорт Excel','',entries,clean(p['reason']));self._finish(c)
        return {'count':len(entries)}

    def _quality(self,c):
        assets=self._all(c,'assets');people={e['id']:e for e in self._all(c,'employees')}
        objects={o[0]:o for o in self._meta(c,'objects') or []};serials={}
        for a in assets:
            serial=clean(a['serial']).casefold()
            if serial:serials.setdefault(serial,[]).append(a['id'])
        issues=[]
        def add(a,message):issues.append({'id':a['id'],'message':message})
        for a in assets:
            if not clean(a['serial']):add(a,'Не указан серийный номер — используйте ID оборудования для идентификации')
            elif len(serials[clean(a['serial']).casefold()])>1:add(a,'Повтор серийного номера: ID '+', '.join(map(str,serials[clean(a['serial']).casefold()])))
            if not clean(a['kit']):add(a,'Не заполнена комплектность')
            if not clean(a['condition']):add(a,'Не заполнено состояние')
            kind=a['target_kind'];target=a['target_id']
            if kind=='Требует уточнения':add(a,'Уточните получателя: неоднозначное имя')
            if kind=='Сотрудник':
                if target not in people:add(a,'Получатель отсутствует в списке сотрудников')
                elif people[target]['status']=='Уволен':add(a,'Получатель уволен')
                elif a['owner']!=people[target]['name']:add(a,'Имя получателя не соответствует выбранному сотруднику')
            elif kind=='Объект' and target not in objects:add(a,'Получатель отсутствует в списке объектов')
            elif kind=='Объект' and a['owner']!=objects[target][1]:add(a,'Название получателя не соответствует объекту')
            if a['status']=='В эксплуатации' and not target:add(a,'Не указан получатель')
            if a['status']=='На складе' and (target or clean(a['owner'])):add(a,'На складе, но закреплено за получателем')
            if a['status']=='Списано' and (target or clean(a['owner'])):add(a,'Списанное оборудование закреплено за получателем')
            if a['status'] not in ['В эксплуатации','На складе','Поломка','В ремонте','Списано']:add(a,'Неизвестный статус')
            if a['status']!='Списано':
                r=people.get(a['responsible_id'])
                if a['responsible_id'] and not r:add(a,'Ответственный отсутствует в списке сотрудников')
                elif r and r['status']=='Уволен':add(a,'Ответственный уволен — назначьте нового')
                elif r and a['responsible']!=r['name']:add(a,'Имя ответственного не соответствует выбранному сотруднику')
                elif not r and clean(a['responsible']):add(a,'Уточните ответственного: выберите сотрудника')
                elif not r and a['status'] in ['В эксплуатации','Поломка','В ремонте']:add(a,'Не указан ответственный')
        missing=[]
        for row in c.execute('SELECT id,file FROM acts'):
            path=(self.root/row['file']).resolve()
            if not path.is_relative_to(self.root.resolve()) or not path.is_file():missing.append({'id':row['id'],'message':'Не найден файл акта'})
        return {'assets':issues,'documents':missing}

    def movement(self,p):
        with self.lock,contextlib.closing(self._connect()) as c:
            self._revision(c,p);op=p.get('operation')
            require(op in ['Передача','Перемещение'],'Неизвестная операция.')
            require(clean(p.get('reason')),'Укажите основание операции.')
            ids=p.get('ids',[]);require(ids and len(ids)<=100 and len(ids)==len(set(ids)),'Выберите от 1 до 100 разных позиций.')
            location=clean(p.get('location'));require(location in self._meta(c,'locations'),'Выберите расположение.')
            if op=='Передача':
                kind,target,name,person=self._recipient(c,p)
                require(person is None or person['status']=='Действующий','Получатель уволен.')
                responsible=self._get(c,'employees',p.get('responsible_id'))
                require(responsible['status']=='Действующий','Выберите действующего ответственного.')
            entries=[]
            for id in ids:
                a=self._get(c,'assets',id);old=copy.deepcopy(a)
                require(a['status'] not in ['Списано','В ремонте'],'Списанную технику или технику в ремонте перемещать нельзя.')
                if op=='Передача':
                    require(a['target_id'] and a['target_kind'] in ['Сотрудник','Объект'],'Для выдачи со склада используйте команду «Выдать»; сначала уточните получателя.')
                    require(a['status'] in ['В эксплуатации','Поломка'],'Техника не закреплена за получателем.')
                    require((kind,target)!=(a['target_kind'],a['target_id']),'Выбран прежний получатель. Для смены расположения используйте перемещение.')
                    a.update(target_kind=kind,target_id=target,owner=name,responsible_id=responsible['id'],responsible=responsible['name'],location=location)
                else:
                    require(a['location']!=location,'Выбрано текущее расположение.')
                    if a['status']=='На складе':require(not a['target_id'] and not clean(a['owner']),'Складская запись имеет противоречивого получателя; сначала исправьте учет.')
                    a['location']=location
                a['updated']=now();entries.append({'action':op,'before':old,'after':a,'note':clean(p.get('reason'))})
            self._backup(c);path=None
            try:
                c.execute('BEGIN IMMEDIATE')
                path=self._act(c,op,name if op=='Передача' else location,entries,{'basis':p['reason']})
                for entry in entries:self._put(c,'assets',entry['after'])
                self._log(c,op,name if op=='Передача' else location,entries,clean(p['reason']));self._finish(c);c.commit()
            except Exception:
                c.rollback()
                if path:(self.root/path).unlink(missing_ok=True)
                raise
            return {'file':path,'count':len(entries)}

    def repair(self,p):
        with self.lock,contextlib.closing(self._connect()) as c:
            self._revision(c,p);a=self._get(c,'assets',p['id']);before=copy.deepcopy(a)
            require(clean(p.get('reason')),'Укажите причину / результат ремонта.')
            jobs=self._meta(c,'repairs') or []
            active=next((r for r in jobs if r['asset_id']==a['id'] and r['state']=='Открыт'),None)
            phase=p.get('phase');require(phase in ['start','finish'],'Неизвестная операция ремонта.')
            if phase=='start':
                require(a['status'] not in ['Списано','В ремонте'] and not active,'Техника списана или уже в ремонте.')
                require(clean(p.get('provider')),'Укажите сервис или исполнителя ремонта.')
                job={'id':max([r['id'] for r in jobs]+[0])+1,'asset_id':a['id'],'state':'Открыт','started':now(),'provider':clean(p['provider']),'reason':clean(p['reason']),'before':before}
                jobs.append(job);a.update(status='В ремонте',location='Сервис: '+job['provider']);op='Передача в ремонт'
            else:
                require(active is not None and a['status']=='В ремонте','Открытый ремонт не найден.')
                require(p.get('result') in ['Исправно','Неисправно'],'Выберите результат ремонта.')
                restored_location=clean(p.get('location'));require(restored_location in self._meta(c,'locations'),'Выберите расположение после ремонта.')
                require(clean(p.get('condition')),'Укажите состояние после ремонта.')
                active.update(state='Завершен',finished=now(),result=p['result'],resolution=clean(p['reason']))
                a.update(status=('В эксплуатации' if a['target_id'] else 'На складе') if p['result']=='Исправно' else 'Поломка',location=restored_location,condition=clean(p['condition']));op='Завершение ремонта'
            a['updated']=now();entries=[{'action':op,'before':before,'after':a,'note':clean(p['reason'])}]
            self._backup(c);path=None
            try:
                c.execute('BEGIN IMMEDIATE');path=self._act(c,op,a['owner'] or 'Склад',entries,{'basis':p['reason']})
                self._put(c,'assets',a);self._meta(c,'repairs',jobs);self._log(c,op,a['name'],entries,clean(p['reason']));self._finish(c);c.commit()
            except Exception:
                c.rollback()
                if path:(self.root/path).unlink(missing_ok=True)
                raise
            return {'file':path}

    def restore_preview(self,file):
        file=Path(file)
        require(file.is_file(),'Архив не найден.')
        with zipfile.ZipFile(file) as z:
            names=z.namelist();require(len(names)==len(set(names)),'В архиве повторяются пути.')
            require(sum(i.file_size for i in z.infolist())<=512*1024*1024,'Архив слишком большой (максимум 512 МБ после распаковки).')
            require('manifest.json' in names,'Это не полный архив Матучета.')
            manifest=json.loads(z.read('manifest.json'))
            require(manifest.get('format')=='matuchet-backup-v1','Неподдерживаемый формат архива.')
            files=manifest.get('files',{});require(set(names)==set(files)|{'manifest.json'} and 'Учет.sqlite' in files,'Неполный состав архива.')
            for name,digest in files.items():
                parts=Path(name).parts
                require('\\' not in name and not name.startswith('/') and '..' not in parts and ':' not in name,'Недопустимый путь в архиве.')
                require(name in ['Учет.sqlite','Шаблон_акта.docx','Шаблон_приема_передачи.docx'] or (name.startswith('Акты/') and name.endswith('.docx')),'Неподдерживаемый файл в архиве.')
                require(hashlib.sha256(z.read(name)).hexdigest()==digest,'Поврежден файл: '+name)
            with tempfile.TemporaryDirectory(prefix='matuchet-check-',dir=self.root) as tmp:
                db=Path(tmp)/'check.sqlite';db.write_bytes(z.read('Учет.sqlite'))
                with contextlib.closing(sqlite3.connect(db)) as con:
                    require(con.execute('PRAGMA integrity_check').fetchone()[0]=='ok','База в архиве повреждена.')
                    tables={r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
                    require({'meta','assets','employees','acts','history'}<=tables,'Архив содержит другую базу.')
                    require(con.execute("SELECT data FROM meta WHERE key='seeded'").fetchone() is not None,'База не инициализирована.')
                    for row in con.execute('SELECT file FROM acts'):
                        require(row[0] in files,'В архиве отсутствует акт: '+str(row[0]))
            return {'file':str(file),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'created':manifest['created'],'revision':manifest['revision'],'documents':sum(n.startswith('Акты/') for n in files)}

    def restore(self,p,*,_shared=False,_safety=None,_local_settings=None):
        with self.lock:
            with contextlib.closing(self._connect()) as c:self._revision(c,p)
            plan=self.restore_preview(p['file']);require(plan['sha256']==p['sha256'],'Архив изменился. Выберите его заново.')
            if _shared and _safety:
                self.restore_preview(_safety)
                safety=_safety
            else:safety=self.backup()['file']
            local_settings=_local_settings if _local_settings is not None else self.snapshot()['settings']
            with tempfile.TemporaryDirectory(prefix='restore-',dir=self.root) as tmp:
                stage=Path(tmp);new=stage/'new';old=stage/'old';new.mkdir();old.mkdir();(new/'Акты').mkdir()
                with zipfile.ZipFile(p['file']) as z:
                    for name in json.loads(z.read('manifest.json'))['files']:
                        dest=new/name;dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes(z.read(name))
                moved=[];installed=[]
                try:
                    for name in ['Учет.sqlite','Акты','Шаблон_акта.docx','Шаблон_приема_передачи.docx']:
                        if (self.root/name).exists():os.replace(self.root/name,old/name);moved.append(name)
                        if (new/name).exists():os.replace(new/name,self.root/name);installed.append(name)
                    with contextlib.closing(self._connect()) as c:
                        c.execute('BEGIN IMMEDIATE')
                        if _shared:
                            settings=self._meta(c,'settings')
                            for key in ['operator','role']:settings[key]=local_settings.get(key,'')
                            self._meta(c,'settings',settings)
                        else:
                            self._meta(c,'revision',max(p['revision'],plan['revision'])+1)
                            self._log(c,'Восстановление','',[],'Восстановлен архив '+plan['created'])
                        c.commit()
                except Exception:
                    for name in installed:os.replace(self.root/name,new/name)
                    for name in moved:os.replace(old/name,self.root/name)
                    raise
            return {'safety_backup':safety}


    def admin_asset(self,p):
        """Explicit record correction; no movement document is generated."""
        with self.lock,self._connect() as c:
            self._revision(c,p);require(clean(p.get('reason')),'Укажите причину редактирования.')
            creating=not p.get('id')
            old={k:'' for k in FIELDS} if creating else self._get(c,'assets',p['id'])
            if creating:old.update(id=self._meta(c,'next_asset'),status='На складе',location='Склад')
            if creating:require(not c.execute('SELECT 1 FROM assets WHERE id=?',(old['id'],)).fetchone(),'Счетчик ID поврежден. Проверьте базу перед добавлением.')
            a=copy.deepcopy(old)
            for k in ['type','name','serial','status','kit','notes','location','purpose','condition','purchase_date','supply_doc']:
                if k in p:a[k]=clean(p[k])
            if 'purchase_price' in p:a['purchase_price']=price(p['purchase_price'])
            require(a['name'],'Название обязательно.')
            require(a['type'] in self._meta(c,'types'),'Выберите тип оборудования.')
            require(a['status'] in ['На складе','В эксплуатации','Поломка','В ремонте','Списано'],'Выберите статус.')
            require(not a['purchase_date'] or self._valid_date(a['purchase_date']),'Дата покупки должна иметь вид ГГГГ-ММ-ДД.')
            require(not a['serial'] or not any(x['id']!=a['id'] and clean(x['serial']).casefold()==a['serial'].casefold() for x in self._all(c,'assets')),'Серийный номер уже существует.')
            if 'target_kind' in p:
                if p.get('target_id'):
                    kind,id,name,_=self._recipient(c,p);a.update(target_kind=kind,target_id=id,owner=name)
                else:a.update(target_kind='',target_id='',owner='')
            if 'responsible_id' in p:
                if p.get('responsible_id'):
                    e=self._get(c,'employees',p['responsible_id']);a.update(responsible_id=e['id'],responsible=e['name'])
                else:a.update(responsible_id='',responsible='')
            if a['status'] in ['На складе','Списано']:
                require(not a['target_id'] and not clean(a['owner']),'Для складской или списанной техники выберите «Без получателя».')
            if a['status']=='В эксплуатации':require(bool(a['target_id']),'Для эксплуатации выберите получателя.')
            jobs=self._meta(c,'repairs') or [];active=next((j for j in jobs if j['asset_id']==a['id'] and j['state']=='Открыт'),None)
            if a['status']=='В ремонте' and not active:
                require(clean(p.get('repair_provider')),'Укажите сервис / исполнителя ремонта.')
                jobs.append({'id':max([j['id'] for j in jobs]+[0])+1,'asset_id':a['id'],'state':'Открыт','started':now(),'provider':clean(p['repair_provider']),'reason':clean(p['reason']),'before':old,'administrative':True})
            elif a['status']!='В ремонте' and active:
                active.update(state='Завершен',finished=now(),result='Административное исправление',resolution=clean(p['reason']))
            elif active and clean(p.get('repair_provider')):active['provider']=clean(p['repair_provider'])
            a.update(updated=now(),warnings='')
            self._backup(c)
            self._put(c,'assets',a);self._meta(c,'repairs',jobs)
            locations=self._meta(c,'locations')
            if a['location'] and a['location'] not in locations:locations.append(a['location']);self._meta(c,'locations',locations)
            if creating:self._meta(c,'next_asset',a['id']+1)
            self._log(c,'Добавление без акта' if creating else 'Редактирование без акта',a['name'],[{'before':None if creating else old,'after':a}],clean(p['reason']));self._finish(c)
        return {'id':a['id']}


    @staticmethod
    def _valid_date(value):
        try:return dt.date.fromisoformat(value).isoformat()==value
        except ValueError:return False


    def admin_employee(self,p):
        with self.lock,self._connect() as c:
            self._revision(c,p);require(clean(p.get('reason')),'Укажите причину редактирования.')
            e=self._get(c,'employees',p['id']);old=copy.deepcopy(e)
            e.update(name=clean(p.get('name')),role=clean(p.get('role')),status=clean(p.get('status')))
            require(e['name'] and e['role'],'Заполните ФИО и должность.')
            require(e['status'] in ['Действующий','Уволен'],'Выберите статус сотрудника.')
            require(not any(x['id']!=e['id'] and clean(x['name']).casefold()==e['name'].casefold() for x in self._all(c,'employees')),'Сотрудник с таким ФИО уже существует.')
            self._backup(c);entries=[{'before':old,'after':e}];self._put(c,'employees',e)
            for a in self._all(c,'assets'):
                before=copy.deepcopy(a)
                if a['target_kind']=='Сотрудник' and a['target_id']==e['id']:a['owner']=e['name']
                if a['responsible_id']==e['id']:a['responsible']=e['name']
                if a!=before:
                    a['updated']=now();self._put(c,'assets',a);entries.append({'before':before,'after':a})
            self._log(c,'Редактирование сотрудника без акта',e['name'],entries,clean(p['reason']));self._finish(c)
        return {}


    def admin_dictionary(self,p):
        with self.lock,self._connect() as c:
            self._revision(c,p);key=p.get('key');require(key in ['types','locations','objects'],'Неизвестный справочник.')
            require(clean(p.get('reason')),'Укажите причину редактирования.')
            value=clean(p.get('value'));require(value,'Введите новое название.')
            items=self._meta(c,key);before=copy.deepcopy(items)
            if key=='objects':
                obj=next((o for o in items if o[0]==p.get('id')),None);require(obj is not None,'Объект не найден.');oldname=obj[1]
                require(not any(o[0]!=obj[0] and o[1].casefold()==value.casefold() for o in items),'Такое название уже существует.')
                obj[1]=value
                if 'kind' in p:obj[2]=clean(p['kind'])
            else:
                oldname=p.get('old');require(oldname in items,'Значение не найдено.')
                require(not any(x!=oldname and x.casefold()==value.casefold() for x in items),'Такое название уже существует.')
                items[items.index(oldname)]=value
            self._backup(c);entries=[{'before':{'name':oldname,'dictionary_values':before},'after':{'name':value,'dictionary_values':copy.deepcopy(items)},'entity':key}];self._meta(c,key,items)
            for a in self._all(c,'assets'):
                old=copy.deepcopy(a)
                if key=='objects' and a['target_kind']=='Объект' and a['target_id']==p.get('id'):a['owner']=value
                if key=='types' and a['type']==oldname:a['type']=value
                if key=='locations' and a['location']==oldname:a['location']=value
                if a!=old:a['updated']=now();self._put(c,'assets',a);entries.append({'before':old,'after':a})
            if key=='locations':
                jobs=self._meta(c,'repairs') or []
                for j in jobs:
                    if j['state']=='Открыт' and j['before'].get('location')==oldname:j['before']['location']=value
                self._meta(c,'repairs',jobs)
            detail=clean(p['reason'])+'; '+oldname+' → '+value
            self._log(c,'Редактирование справочника без акта',value,entries,detail);self._finish(c)
        return {}
